import io

import fitz  # PyMuPDF
from docx import Document

class ParseError(Exception):
    pass

def extract_text(filename: str, data: bytes) -> str:
    ext = filename.lower().rsplit(".", 1)[-1]
    try:
        text = _pdf(data) if ext == "pdf" else _docx(data)
    except ParseError:
        raise
    except Exception as e:  # corrupt / unsupported internals
        raise ParseError("The file could not be read. It may be corrupted or password-protected.") from e
    text = "\n".join(l.rstrip() for l in text.splitlines()).strip()
    if len(text) < 20:
        raise ParseError("No readable text found. Scanned/image-only resumes are not supported; export a text-based PDF or DOCX.")
    return text

def _pdf(data: bytes) -> str:
    with fitz.open(stream=data, filetype="pdf") as doc:
        if doc.needs_pass:
            raise ParseError("This PDF is password-protected.")
        return "\n".join(p.get_text() for p in list(doc)[:20])

def _docx(data: bytes) -> str:
    d = Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append("  ".join(c.text for c in row.cells))
    return "\n".join(parts)
